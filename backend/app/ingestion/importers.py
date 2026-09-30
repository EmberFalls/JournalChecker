import csv
import json
from datetime import UTC, datetime
from pathlib import Path
from sqlalchemy.orm import Session
from app.core.enums import EvidenceState
from app.core.normalization import normalize_issn
from app.db.models import Evidence, Journal, JournalIdentifier, Provider

DOAJ_DATASET_URL = "https://doaj.org/csv"
SCOPUS_SOURCE_LIST_PAGE = "https://www.elsevier.com/products/scopus/content"


def _provider(db: Session, name: str) -> Provider:
    item = db.query(Provider).filter_by(name=name).one_or_none()
    if not item:
        item = Provider(name=name)
        db.add(item)
    item.last_imported_at = datetime.now(UTC)
    return item


def _upsert_journal(db: Session, title: str, issns: list[str], publisher: str | None, domain: str | None = None) -> Journal | None:
    found_rows = [
        db.query(JournalIdentifier).filter_by(scheme="ISSN", value=issn).one_or_none()
        for issn in issns
    ]
    existing_ids = {item.journal_id for item in found_rows if item}
    if len(existing_ids) == 1:
        journal = db.get(Journal, next(iter(existing_ids)))
        known = {item.value for item in journal.identifiers if item.scheme == "ISSN"}
        for issn in issns:
            if issn not in known and not db.query(JournalIdentifier).filter_by(scheme="ISSN", value=issn).first():
                db.add(JournalIdentifier(journal_id=journal.id, scheme="ISSN", value=issn))
        if not journal.canonical_publisher and publisher:
            journal.canonical_publisher = publisher
        if not journal.official_domain and domain:
            journal.official_domain = domain
        db.flush()
        return journal
    if len(existing_ids) > 1:
        # Conflicting source records already associate these identifiers to distinct
        # journals. Never silently merge identities during ingestion.
        return None
    journal = Journal(current_title=title, canonical_publisher=publisher, official_domain=domain)
    db.add(journal)
    db.flush()
    for issn in issns:
        db.add(JournalIdentifier(journal_id=journal.id, scheme="ISSN", value=issn))
    return journal


def _add_evidence_once(db: Session, *, journal_id: str, provider: str, evidence_type: str, value: dict, source_url: str | None, confidence: float = 0.95) -> None:
    """Keep repeated imports of the same source record idempotent."""
    existing = db.query(Evidence).filter_by(
        journal_id=journal_id,
        provider=provider,
        evidence_type=evidence_type,
        source_url=source_url,
    ).one_or_none()
    if existing:
        existing.value = value
        existing.observed_at = datetime.now(UTC)
    else:
        db.add(Evidence(journal_id=journal_id, provider=provider, evidence_type=evidence_type, value=value, status=EvidenceState.VERIFIED, confidence=confidence, source_url=source_url))


def _domain(url: str | None) -> str | None:
    if not url:
        return None
    from urllib.parse import urlparse
    return urlparse(url).hostname


def import_doaj_csv(db: Session, path: str, source_url: str = DOAJ_DATASET_URL) -> int:
    """Import the official public DOAJ journal CSV. Absence is never stored as a negative finding."""
    _provider(db, "doaj")
    imported = 0
    with Path(path).open(encoding="utf-8-sig", newline="") as file:
        for row in csv.DictReader(file):
            title = row.get("Journal title")
            issns = [normalize_issn(row.get(key, "")) for key in ("Journal ISSN (print version)", "Journal EISSN (online version)")]
            issns = [value for value in issns if value]
            if not title or not issns:
                continue
            journal = _upsert_journal(db, title, issns, row.get("Publisher"), _domain(row.get("Journal URL")))
            if journal is None:
                continue
            _add_evidence_once(db, journal_id=journal.id, provider="doaj", evidence_type="doaj_listing", source_url=row.get("URL in DOAJ") or source_url, value={
                "title": title, "publisher": row.get("Publisher"), "journal_url": row.get("Journal URL"),
                "review_process": row.get("Review process"), "license": row.get("Journal license"),
                "last_updated": row.get("Last updated Date"), "dataset_url": source_url,
            })
            imported += 1
    db.commit()
    return imported


def import_doaj_json(db: Session, path: str) -> int:
    records = json.loads(Path(path).read_text(encoding="utf-8"))
    _provider(db, "doaj")
    imported = 0
    for item in records:
        bib = item.get("bibjson", item)
        issns = [normalize_issn(x) for x in bib.get("pissn", "").split(",") + bib.get("eissn", "").split(",")]
        issns = [x for x in issns if x]
        if not issns or not bib.get("title"):
            continue
        journal = _upsert_journal(db, bib["title"], issns, bib.get("publisher"))
        if journal is None:
            continue
        _add_evidence_once(db, journal_id=journal.id, provider="doaj", evidence_type="doaj_listing", value={"title": bib["title"]}, source_url=item.get("id"))
        imported += 1
    db.commit()
    return imported


def import_scopus_csv(db: Session, path: str) -> int:
    _provider(db, "scopus")
    imported = 0
    with Path(path).open(encoding="utf-8-sig", newline="") as file:
        for row in csv.DictReader(file):
            title = row.get("Source Title") or row.get("Title")
            issns = [normalize_issn(row.get(key, "")) for key in ("Print-ISSN", "E-ISSN", "ISSN")]
            issns = [x for x in issns if x]
            if not title or not issns:
                continue
            journal = _upsert_journal(db, title, issns, row.get("Publisher's Name") or row.get("Publisher"))
            if journal is None:
                continue
            active = (row.get("Active or Inactive") or row.get("Status") or "").lower() != "inactive"
            _add_evidence_once(db, journal_id=journal.id, provider="scopus", evidence_type="scopus_coverage", value={"coverage": row.get("Coverage"), "active": active, "source_id": row.get("Source ID")}, source_url=f"scopus-source:{row.get('Source ID', title)}")
            imported += 1
    db.commit()
    return imported


def import_scopus_xlsx(db: Session, path: str, source_url: str = SCOPUS_SOURCE_LIST_PAGE) -> int:
    """Import the official Scopus Source Title List XLSX without querying or scraping Scopus."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependency declaration is tested in deployment
        raise RuntimeError("openpyxl is required to import Scopus XLSX files") from exc
    _provider(db, "scopus")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    imported = 0
    for values in rows:
        row = dict(zip(headers, values))
        title = row.get("Source Title")
        issns = [normalize_issn(str(row.get(key) or "")) for key in ("ISSN", "EISSN")]
        issns = [value for value in issns if value]
        if not title or not issns:
            continue
        journal = _upsert_journal(db, str(title), issns, row.get("Publisher") or row.get("Publisher Imprints Grouped to Main Publisher"))
        if journal is None:
            continue
        active = str(row.get("Active or Inactive") or "").strip().lower() == "active"
        source_id = str(row.get("Sourcerecord ID") or title)
        _add_evidence_once(db, journal_id=journal.id, provider="scopus", evidence_type="scopus_coverage", source_url=f"{source_url}#source-{source_id}", value={
            "title": title, "source_id": source_id, "active": active, "coverage": row.get("Coverage"),
            "source_type": row.get("Source Type"), "open_access_status": row.get("Open Access Status"),
            "dataset_url": source_url,
        })
        imported += 1
    db.commit()
    return imported


def import_scimago_csv(db: Session, path: str, source_url: str = "https://www.scimagojr.com/journalrank.php") -> int:
    """Import a user-downloaded SCImago ranking CSV; does not fetch or scrape the site."""
    _provider(db, "scimago")
    imported = 0
    with Path(path).open(encoding="utf-8-sig", newline="") as file:
        sample = file.read(4096)
        file.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(file, dialect=dialect)
        for row in reader:
            normalized_headers = {str(key).strip().lower(): key for key in (row or {}) if key}

            def field(*names: str) -> str | None:
                key = next((normalized_headers[name.lower()] for name in names if name.lower() in normalized_headers), None)
                value = row.get(key) if key else None
                return str(value).strip() if value not in (None, "") else None

            title = field("Title", "Journal title", "Source title")
            raw_issns = field("Issn", "ISSN", "ISSN/eISSN")
            issns = [normalize_issn(part.strip()) for part in (raw_issns or "").split(",")]
            issns = [value for value in issns if value]
            if not title or not issns:
                continue
            journal = _upsert_journal(db, title, issns, field("Publisher"))
            if journal is None:
                continue
            year = next((name.split("(")[1].split(")")[0] for name in normalized_headers if name.startswith("sjr (") and "(" in name), None)
            metric_columns = ("SJR", "SJR Best Quartile", "H index", "Type", "Country")
            values = {column: field(column, f"{column} ({year})" if year else column) for column in metric_columns}
            values.update({
                "title": title,
                "issns": issns,
                "year": year,
                "source_id": field("Sourceid", "Source ID", "Source Id"),
                "dataset_url": source_url,
            })
            _add_evidence_once(
                db,
                journal_id=journal.id,
                provider="scimago",
                evidence_type="scimago_metrics",
                value=values,
                source_url=source_url,
            )
            imported += 1
    db.commit()
    return imported
